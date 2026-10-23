"""Turn sales scenarios into non-factual response-structure guidance."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


_INTENTS = {
    "hỏi giá": "price", "chê giá cao": "price", "hỏi khuyến mãi": "price",
    "phân vân chọn tour": "recommend_tour", "đi cùng người lớn tuổi": "recommend_tour",
    "đi cùng trẻ nhỏ": "child_price", "đi một mình": "recommend_tour", "khách đoàn": "general",
    "lịch trình & mua sắm": "itinerary", "thời tiết": "general", "đặt cọc & thanh toán": "payment",
    "trả góp": "payment", "giấy tờ": "general", "visa": "visa", "hủy / đổi tour": "cancellation",
    "bảo hiểm": "general", "hành lý": "baggage", "khách chần chừ": "general",
    "hết chỗ": "availability", "khiếu nại": "complaint",
}


def build_sales_guidance(path: Path) -> list[dict]:
    """Read scenario IDs and intent groups only; never copy spreadsheet answers or claims."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        result = []
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not row or not isinstance(row[0], (int, float)) or not isinstance(row[1], str):
                continue
            intent = _INTENTS.get(row[1].strip().casefold(), "general")
            result.append({
                "scenario_id": int(row[0]),
                "intent": intent,
                "response_structure": "Acknowledge the question, establish the tour and relevant context, then answer from its official document.",
                "required_action": "Ask a concise clarifying question when the tour or needed detail is missing.",
                "prohibited_actions": ["Do not state unverified prices, dates, availability, discounts, or policy terms."],
            })
        return result
    finally:
        workbook.close()
